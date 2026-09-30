from __future__ import annotations

from dataclasses import replace
from io import BytesIO
from uuid import uuid4

import pytest
from docx import Document

from app.import_document_preparation import (
    ImportDocumentPreparationError, prepare_import_document,
)
from app.postgres_vector_generations import VectorScope
from hello_agents.memory.embedding import SimpleEmbedding
from hello_agents.memory.rag.embedding_runtime import (
    RAGEmbeddingRuntime, build_rag_embedding,
)
from hello_agents.memory.rag.index_identity import IndexIdentity
from hello_agents.memory.rag.prepare import qdrant_point_id
from hello_agents.memory.rag.qdrant_pipeline import split_qdrant_text


class _FakeEmbedding(SimpleEmbedding):
    def __init__(self):
        super().__init__(dimension=384)
        self.calls: list[list[str]] = []

    def encode(self, texts):
        self.calls.append(list(texts))
        return [[1.0, 2.0] + [0.0] * 382 for _ in texts]


def _context():
    task_id, user_id, document_id = (str(uuid4()) for _ in range(3))
    profile = build_rag_embedding({}, backend="qdrant").profile
    scope = VectorScope(user_id, "rag", f"pdf_{user_id}",
                        IndexIdentity("qdrant", "documents", profile))
    engine = _FakeEmbedding()
    runtime = RAGEmbeddingRuntime(profile, engine)
    return dict(task_id=task_id, user_id=user_id, document_id=document_id,
                scope=scope, embedding_runtime=runtime), engine


def _prepare(source: bytes, name="paper.txt", suffix=".txt", **changes):
    kwargs, engine = _context()
    kwargs.update(source_bytes=source, original_name=name, suffix=suffix)
    kwargs.update(changes)
    return prepare_import_document(**kwargs), kwargs, engine


def _pdf_bytes(*page_texts: str) -> bytes:
    """Build a small real PDF with extractable text and a page tree."""
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"",
               b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    page_ids = []
    for text in page_texts:
        page_id = len(objects) + 1
        content_id = page_id + 1
        page_ids.append(page_id)
        objects.append((f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
                        f"/Resources << /Font << /F1 3 0 R >> >> "
                        f"/Contents {content_id} 0 R >>").encode())
        stream = f"BT /F1 12 Tf 72 700 Td ({text}) Tj ET".encode("ascii")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode()
                       + stream + b"\nendstream")
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = (f"<< /Type /Pages /Kids [{kids}] "
                  f"/Count {len(page_ids)} >>").encode()
    data = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(data))
        data.extend(f"{number} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref = len(data)
    data.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for offset in offsets[1:]:
        data.extend(f"{offset:010d} 00000 n \n".encode())
    data.extend((f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\n"
                 f"startxref\n{xref}\n%%EOF\n").encode())
    return bytes(data)


def test_pdf_bytes_preserve_real_page_numbers_and_original_name():
    points, kwargs, engine = _prepare(_pdf_bytes("alpha", "", "gamma"),
                                      name="paper.pdf", suffix=".pdf")
    assert [point.payload["metadata"]["page_number"] for point in points] == [1, 3]
    assert [point.payload["content"] for point in points] == ["alpha", "gamma"]
    assert all(point.payload["metadata"]["file_name"] == "paper.pdf"
               for point in points)
    assert all("file_path" not in point.payload["metadata"] for point in points)
    assert engine.calls == [["alpha", "gamma"]]


def test_docx_bytes_join_nonempty_paragraphs_before_qdrant_window():
    document = Document()
    document.add_paragraph("x" * 790)
    document.add_paragraph("   ")
    document.add_paragraph("y" * 40)
    stream = BytesIO()
    document.save(stream)
    points, _, _ = _prepare(stream.getvalue(), "notes.docx", ".docx")
    assert [point.payload["content"] for point in points] == split_qdrant_text(
        "x" * 790 + "\n\n" + "y" * 40)


@pytest.mark.parametrize("name,suffix", [("note.txt", ".txt"),
                                          ("note.md", ".md"),
                                          ("note.markdown", ".markdown")])
def test_text_decodes_bom_then_gbk_and_markdown_alias(name, suffix):
    utf8, _, _ = _prepare(b"\xef\xbb\xbfhello", name, suffix)
    gbk, _, _ = _prepare("中文".encode("gbk"), name, suffix)
    assert utf8[0].payload["content"] == "hello"
    assert gbk[0].payload["content"] == "中文"


def test_qdrant_window_and_stable_document_isolated_ids():
    content = ("x" * 800 + "y" * 100).encode()
    points, kwargs, _ = _prepare(content)
    assert [point.payload["content"] for point in points] == [
        "x" * 800, "x" * 120 + "y" * 100,
    ]
    assert [point.id for point in points] == [
        qdrant_point_id(kwargs["scope"].namespace, kwargs["document_id"], index)
        for index in (0, 1)
    ]
    repeated = prepare_import_document(**kwargs)
    assert [point.id for point in repeated] == [point.id for point in points]
    changed = dict(kwargs, document_id=str(uuid4()))
    assert prepare_import_document(**changed)[0].id != points[0].id
    assert split_qdrant_text(" " + "x" * 800 + "y" * 100 + " ") == [
        point.payload["content"] for point in points
    ]
    for point in points:
        payload = point.payload
        assert payload["document_id"] == kwargs["document_id"]
        assert payload["rag_namespace"] == kwargs["scope"].namespace
        assert payload["metadata"]["embedding_fingerprint"] == (
            kwargs["scope"].identity.profile.fingerprint)
        assert len(point.vector) == 384


@pytest.mark.parametrize("source,name,suffix", [
    (b"  \n", "blank.txt", ".txt"),
    (_pdf_bytes(""), "blank.pdf", ".pdf"),
    (b"bad pdf", "invalid.pdf", ".pdf"),
    (b"\x81", "invalid.txt", ".txt"),
    (b"z", "wrong.exe", ".exe"),
    (b"z", "wrong.md", ".txt"),
    (b"z", r"C:\upload\wrong.txt", ".txt"),
])
def test_invalid_or_blank_source_returns_no_points(source, name, suffix):
    kwargs, engine = _context()
    with pytest.raises(ImportDocumentPreparationError):
        prepare_import_document(source_bytes=source, original_name=name,
                                suffix=suffix, **kwargs)
    assert engine.calls == []


def test_runtime_mutation_during_embedding_rejects_result():
    kwargs, engine = _context()
    runtime = kwargs["embedding_runtime"]

    def mutating_encode(texts):
        object.__setattr__(runtime, "profile",
                           replace(runtime.profile, revision="changed"))
        return [[1.0, 2.0] + [0.0] * 382 for _ in texts]

    engine.encode = mutating_encode
    with pytest.raises(ImportDocumentPreparationError, match="embedding_identity"):
        prepare_import_document(source_bytes=b"hello", original_name="note.txt",
                                suffix=".txt", **kwargs)


def test_wrong_scope_profile_and_size_fail_before_embedding():
    kwargs, engine = _context()
    wrong = replace(kwargs["scope"].identity.profile, revision="other")
    wrong_scope = replace(kwargs["scope"], identity=IndexIdentity("qdrant", "documents", wrong))
    for change in (
        {"scope": wrong_scope},
        {"scope": replace(kwargs["scope"], namespace="pdf_other")},
        {"scope": replace(kwargs["scope"], tenant_id=str(uuid4()))},
        {"document_id": "not-a-uuid"},
    ):
        with pytest.raises(ImportDocumentPreparationError):
            prepare_import_document(source_bytes=b"hello", original_name="note.txt",
                                    suffix=".txt", **(kwargs | change))
    with pytest.raises(ImportDocumentPreparationError, match="source_limit"):
        prepare_import_document(source_bytes=b"z" * (100 * 1024 * 1024 + 1),
                                original_name="note.txt", suffix=".txt", **kwargs)
    assert engine.calls == []
