"""Pure DOCX rendering of a saved learning report snapshot."""
from __future__ import annotations

from datetime import datetime
from io import BytesIO

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn


def render_report_docx(report: str, user_id: str, generated_at: datetime | None = None) -> bytes:
    if generated_at is None:
        generated_at = datetime.now()
    doc = Document()

    # =========================
    # 全局字体设置
    # =========================
    styles = doc.styles

    normal_style = styles["Normal"]
    normal_style.font.name = "微软雅黑"
    normal_style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    normal_style.font.size = Pt(11)

    # =========================
    # 标题
    # =========================
    title = doc.add_heading("PDF 智能学习报告", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for run in title.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(20)
        run.bold = True

    # =========================
    # 基本信息
    # =========================
    info = doc.add_paragraph()
    info.alignment = WD_ALIGN_PARAGRAPH.CENTER

    run = info.add_run(
        f"生成时间：{generated_at.strftime('%Y-%m-%d %H:%M:%S')}    "
        f"用户：{user_id}"
    )
    run.font.name = "微软雅黑"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(10)

    doc.add_paragraph("")

    # =========================
    # 正文
    # =========================
    for raw_line in report.splitlines():
        line = raw_line.strip()

        if not line:
            doc.add_paragraph("")
            continue

        # 去掉开头的报告标题，避免重复
        if line.startswith("📘 PDF 智能学习报告"):
            continue

        # 一级小节标题
        if (
                line.startswith("一、")
                or line.startswith("二、")
                or line.startswith("三、")
                or line.startswith("四、")
                or line.startswith("五、")
                or line.startswith("六、")
                or line.startswith("七、")
                or line.startswith("八、")
                or line.startswith("九、")
                or line.startswith("十、")
        ):
            heading = doc.add_heading(line, level=2)

            for run in heading.runs:
                run.font.name = "微软雅黑"
                run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
                run.font.size = Pt(15)
                run.bold = True

            continue

        # 项目符号
        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            run = p.add_run(line[2:])
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            run.font.size = Pt(11)
            continue

        # 数字列表
        if len(line) > 2 and line[0].isdigit() and line[1] in [".", "．", "、"]:
            p = doc.add_paragraph(style="List Number")
            run = p.add_run(line)
            run.font.name = "微软雅黑"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
            run.font.size = Pt(11)
            continue

        # 普通段落
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.25

        run = p.add_run(line)
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(11)

    # =========================
    # 页脚式结尾
    # =========================
    doc.add_paragraph("")
    end = doc.add_paragraph("—— 由 PDF 智能学习助手自动生成 ——")
    end.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for run in end.runs:
        run.font.name = "微软雅黑"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
        run.font.size = Pt(9)
        run.italic = True

    output = BytesIO()
    doc.save(output)

    return output.getvalue()
